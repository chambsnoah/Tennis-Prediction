"""Read-only commissioner workbook preview; never a season-state importer."""

from __future__ import annotations

from contextlib import ExitStack
from dataclasses import dataclass, field, replace
from io import BytesIO
from pathlib import Path
import posixpath
import re
import stat
import warnings
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

from tennis_api.models.canonical import (
    Draw, DrawSlot, EntryKind, Event, PlayerId, Provenance, Tour,
)
from tennis_api.models.identity import IdentityRegistry, IdentityResolutionError


TOURNAMENTS = (
    "Australie", "Indian Wells", "Miami", "Monte Carlo", "Madrid", "Rome",
    "Roland-Garros", "Wimbledon", "Montr\u00e9al", "Cincinnati", "US Open",
    "Shanghai", "Paris",
)
_SLAMS = {"Australie", "Roland-Garros", "Wimbledon", "US Open"}
_SIX_ROUNDS = {"Monte Carlo", "Paris"}
_EVENT_LABELS = {
    "Australian Open": "Australie", "French Open": "Roland-Garros",
    "Canadian Masters": "Montr\u00e9al", "Toronto": "Montr\u00e9al", "Montreal": "Montr\u00e9al",
}
_FORMATS = {"GS128": (128, 128, 32, 7, 5), "M7": (128, 96, 32, 7, 3),
            "M6": (64, 56, 16, 6, 3)}
_SUPPORT = ("Rankings", "Choix", "Classement", "Banques (Solde restant)", "ChoixSommaire")
_BANDS = ("1-4", "5-8", "9-16", "17-32")


class WorkbookImportError(ValueError):
    """Safe reason codes only: no filenames, participant labels or cell values."""


@dataclass(frozen=True)
class WorkbookLimits:
    file_bytes: int = 8 * 1024 * 1024
    zip_bytes: int = 32 * 1024 * 1024
    zip_members: int = 256
    rows: int = 10_000
    columns: int = 128
    cells: int = 300_000

    def __post_init__(self) -> None:
        if any(type(v) is not int or v < 1 for v in vars(self).values()):
            raise ValueError("Invalid workbook limits")


def _bounded_bytes(path: str | Path, limits: WorkbookLimits) -> bytes:
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise WorkbookImportError("unsupported_file_type")
    info = path.stat()
    if not stat.S_ISREG(info.st_mode):
        raise WorkbookImportError("input_must_be_regular_file")
    if info.st_size > limits.file_bytes:
        raise WorkbookImportError("file_byte_limit")
    with path.open("rb") as source:
        data = source.read(limits.file_bytes + 1)
    if len(data) > limits.file_bytes:
        raise WorkbookImportError("file_byte_limit")
    with ZipFile(BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > limits.zip_members:
            raise WorkbookImportError("zip_member_limit")
        if sum(m.file_size for m in members) > limits.zip_bytes:
            raise WorkbookImportError("zip_expanded_byte_limit")
        if len({m.filename for m in members}) != len(members):
            raise WorkbookImportError("duplicate_zip_members")
        cells = dense_cells = 0
        # Inspect *all* members before openpyxl allocates shared strings/styles.
        for member in members:
            if member.flag_bits & 1 or member.compress_type not in (0, 8):
                raise WorkbookImportError("unsupported_zip_member")
            with archive.open(member) as source:
                content = source.read(limits.zip_bytes + 1)
            if len(content) != member.file_size:
                raise WorkbookImportError("zip_expanded_byte_limit")
            if not member.filename.endswith((".xml", ".rels")):
                continue
            # Reject DTD/entity declarations even in UTF-16/32 XML.
            upper = content.replace(b"\x00", b"").upper()
            if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                raise WorkbookImportError("unsafe_xml")
            rows = 0
            worksheet = member.filename.startswith("xl/worksheets/")
            current_row = previous_row = previous_column = max_row = max_column = 0
            for action, node in ElementTree.iterparse(BytesIO(content), events=("start", "end")):
                tag = node.tag.rsplit("}", 1)[-1]
                if action == "start":
                    if worksheet and tag == "row":
                        current_row = int(node.get("r", "0"))
                        if current_row <= previous_row:
                            raise WorkbookImportError("invalid_row_order")
                        previous_row = current_row
                        previous_column = 0
                        max_row = max(max_row, current_row)
                    continue
                if tag == "Relationship" and node.get("Type", "").endswith("/worksheet"):
                    target = node.get("Target", "")
                    base = posixpath.dirname(posixpath.dirname(member.filename))
                    target = posixpath.normpath(
                        target.lstrip("/") if target.startswith("/") else posixpath.join(base, target)
                    )
                    if (node.get("TargetMode") == "External"
                            or not target.startswith("xl/worksheets/") or not target.endswith(".xml")
                            or target not in archive.namelist()):
                        raise WorkbookImportError("unsupported_worksheet_location")
                if worksheet and tag in ("dimension", "c"):
                    refs = node.get("ref", "") if tag == "dimension" else node.get("r", "")
                    for ref in refs.split(":"):
                        match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", ref)
                        if match is None:
                            raise WorkbookImportError("invalid_cell_reference")
                        column = 0
                        for char in match[1]:
                            column = column * 26 + ord(char) - ord("A") + 1
                        if (column > limits.columns or int(match[2]) > limits.rows
                                or column * int(match[2]) > limits.cells):
                            raise WorkbookImportError("worksheet_dimension_limit")
                        if tag == "c":
                            if int(match[2]) != current_row or column <= previous_column:
                                raise WorkbookImportError("invalid_cell_order_or_row")
                            previous_column = column
                            max_column = max(max_column, column)
                if worksheet and tag == "row":
                    rows += 1
                    if rows > limits.rows or int(node.get("r", "0")) > limits.rows:
                        raise WorkbookImportError("worksheet_row_limit")
                if worksheet and tag == "c":
                    cells += 1
                    if cells > limits.cells:
                        raise WorkbookImportError("worksheet_cell_limit")
                node.clear()
            if worksheet:
                dense_cells += max_row * max_column
                if dense_cells > limits.cells:
                    raise WorkbookImportError("worksheet_dense_cell_limit")
    return data


def _sheet_rows(sheet: Any) -> tuple[list[tuple[Any, ...]], set[tuple[int, int]]]:
    # Source dimensions may lie or be absent. Preflight bounds actual coordinates.
    sheet.reset_dimensions()
    rows = []
    formulas: set[tuple[int, int]] = set()
    for i, row in enumerate(sheet.iter_rows()):
        rows.append(tuple(cell.value for cell in row))
        formulas.update((i, j) for j, cell in enumerate(row) if cell.data_type == "f")
    width = max((len(row) for row in rows), default=0)
    return [row + (None,) * (width - len(row)) for row in rows], formulas


@dataclass(frozen=True)
class SheetSummary:
    sheet: str
    records: int
    missing_values: int
    substitutions: int
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookRanking:
    player_id: PlayerId
    rank: int | None
    sheet: str
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookPick:
    participant: str = field(repr=False)
    event: str
    slot: str
    player_id: PlayerId
    band: str | None
    is_substitute: bool
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookStanding:
    participant: str = field(repr=False)
    position: int
    event_points: tuple[int | None, ...]
    bonus: int | None
    total: int | None
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookBank:
    participant: str = field(repr=False)
    slam: tuple[int | None, ...] | None
    masters: tuple[int | None, ...]
    maxed_players: tuple[PlayerId, ...]
    maxed_count: int | None
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookUsage:
    participant: str = field(repr=False)
    player_id: PlayerId
    count: int | None
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookSubstitution:
    participant: str = field(repr=False)
    event: str
    withdrawn_id: PlayerId
    replacement_id: PlayerId
    band: str | None
    provenance: Provenance


@dataclass(frozen=True)
class WorkbookSeason:
    """Private local source observations, not a reconciled or consumed ledger."""

    picks: tuple[WorkbookPick, ...] = field(default=(), repr=False)
    standings: tuple[WorkbookStanding, ...] = field(default=(), repr=False)
    banks: tuple[WorkbookBank, ...] = field(default=(), repr=False)
    usage: tuple[WorkbookUsage, ...] = field(default=(), repr=False)
    substitutions: tuple[WorkbookSubstitution, ...] = field(default=(), repr=False)


@dataclass(frozen=True)
class WorkbookPreview:
    draw: Draw
    rankings: tuple[WorkbookRanking, ...]
    summaries: tuple[SheetSummary, ...]
    missing_sheets: tuple[str, ...]
    unsupported_sheets: tuple[str, ...]
    warnings: tuple[str, ...]
    registry_revision: str
    registry_digest: str
    season: WorkbookSeason = field(repr=False)


def _integer(value: Any, minimum: int = 1, maximum: int = 100_000) -> int:
    if type(value) not in (int, float) or not minimum <= value <= maximum:
        raise ValueError("Invalid workbook integer")
    if int(value) != value:
        raise ValueError("Invalid workbook integer")
    return int(value)


def _rank(value: Any) -> int | None:
    if value in (None, "#N/A"):
        return None
    rank = _integer(value)
    if rank == 9999:
        raise WorkbookImportError("bye_rank_for_real_player")
    return rank


class _UnsupportedLayout(Exception):
    pass


def _support_summary(
    sheet: str, rows: list[tuple[Any, ...]], registry: IdentityRegistry,
    provenance: Provenance,
) -> tuple[SheetSummary, tuple[WorkbookRanking, ...], tuple[str, ...], WorkbookSeason]:
    records = missing = substitutions = 0
    rankings: list[WorkbookRanking] = []
    warnings: list[str] = []
    seen: set[Any] = set()
    picks: list[WorkbookPick] = []
    standings: list[WorkbookStanding] = []
    banks: list[WorkbookBank] = []
    usage: list[WorkbookUsage] = []
    replacements: list[WorkbookSubstitution] = []

    def identity(value: Any) -> PlayerId:
        if not isinstance(value, str) or not value.strip():
            raise WorkbookImportError("invalid_support_player")
        return registry.resolve(source=provenance.source, tour=Tour.ATP, name=value.strip())

    def unique(key: Any) -> None:
        if key in seen:
            raise WorkbookImportError("duplicate_support_record")
        seen.add(key)

    def participant(value: Any) -> str:
        if not isinstance(value, str) or not value.strip():
            raise WorkbookImportError("invalid_participant_label")
        return value.strip()

    def number(value: Any, maximum: int = 100_000, minimum: int = 0) -> int | None:
        nonlocal missing
        if value is None:
            missing += 1
            return None
        else:
            return _integer(value, minimum, maximum)

    def band(value: Any) -> str | None:
        nonlocal missing
        if value is None:
            missing += 1
        elif value not in (*_BANDS, *("#" + b for b in _BANDS), "NF"):
            raise WorkbookImportError("invalid_pick_band")
        return None if value is None else str(value).lstrip("#")

    width = len(rows[0]) if rows else 0
    if sheet == "Rankings":
        if (width < 2 or rows[0][0] not in ("Rank", "Rankings", "Rang", "Classement")
                or rows[0][1] not in ("Player", "Joueur", "Nom")):
            raise _UnsupportedLayout
        for row in rows[1:]:
            if not any(v is not None for v in row):
                continue
            player_id = identity(row[1])
            unique(player_id)
            rank = _rank(row[0])
            missing += rank is None
            rankings.append(WorkbookRanking(player_id, rank, sheet, provenance))
            records += 1
    elif sheet == "Choix":
        if width < 28 or not any(r[0] == "Participants" for r in rows):
            raise _UnsupportedLayout
        current = None
        participants = set()
        for row in rows:
            if row[0] == "Participants":
                # When event headings are supplied, verify their positional order.
                if any(row[2 + i * 2] not in (None, name) for i, name in enumerate(TOURNAMENTS)):
                    raise _UnsupportedLayout
                continue
            if row[0] is not None:
                current = participant(row[0])
                if current in participants:
                    raise WorkbookImportError("duplicate_participant_block")
                participants.add(current)
            if row[1] is None:
                if any(v is not None for v in row[2:]):
                    raise _UnsupportedLayout
                continue
            label = str(row[1])
            alternative = label.startswith(("Substitut", "Remplacant", "Rempla\u00e7ant"))
            if current is None or not (label.startswith("Joueur") or alternative):
                raise _UnsupportedLayout
            for i in range(13):
                value, category = row[2 + i * 2:4 + i * 2]
                if value is None:
                    if category is not None:
                        raise WorkbookImportError("pick_band_without_player")
                    continue
                player_id = identity(value)
                unique((current, i, alternative, player_id))
                category = band(category)
                picks.append(WorkbookPick(current, TOURNAMENTS[i], label, player_id,
                                          category, alternative, provenance))
                records += 1
                substitutions += alternative
    elif sheet == "Classement":
        if (len(rows) < 4 or width < 17 or rows[3][0] not in ("Position", "Rank", "Rang")
                or rows[3][1] not in ("Participants", "Participant", "Nom")
                or rows[3][16] != "Total"
                or any(heading not in (None, name)
                       for heading, name in zip(rows[3][2:15], TOURNAMENTS))):
            raise _UnsupportedLayout
        for row in rows[4:]:
            if not any(v is not None for v in row):
                continue
            who = participant(row[1])
            unique(who)
            position = _integer(row[0])
            points = tuple(number(value) for value in row[2:15])
            bonus, total = number(row[15]), number(row[16])
            standings.append(WorkbookStanding(who, position, points, bonus, total, provenance))
            records += 1
    elif sheet == "Banques (Solde restant)":
        if len(rows) < 2 or width < 6:
            raise _UnsupportedLayout
        masters_only = rows[0][2] == "Tournois Masters"
        if not masters_only and (width < 10 or rows[0][2] != "Tournois Grand Chelem"
                                 or rows[0][6] != "Tournois Masters"):
            raise _UnsupportedLayout
        rows = [row + (None,) * max(0, (9 if masters_only else 13) - width) for row in rows]
        offsets = (2,) if masters_only else (2, 6)
        if any(tuple(str(v).lstrip("#") for v in rows[1][offset:offset + 4]) != _BANDS
               for offset in offsets):
            raise _UnsupportedLayout
        if masters_only:
            warnings.append(sheet + ":slam_quotas_absent")
        for row in rows[2:]:
            if not any(v is not None for v in row):
                continue
            who = participant(row[1])
            unique(who)
            quotas = []
            for offset in offsets:
                maxima = (9, 9, 9, 7) if masters_only or offset == 6 else (4, 4, 4, 4)
                quotas.append(tuple(number(value, maximum)
                                    for value, maximum in zip(row[offset:offset + 4], maxima)))
            count_col, names_col = (7, 8) if masters_only else (11, 12)
            names = row[names_col]
            ids = [] if names is None else [identity(v) for v in str(names).split(",")]
            if len(set(ids)) != len(ids):
                raise WorkbookImportError("duplicate_maxed_player")
            if row[count_col] is None:
                missing += 1
            elif _integer(row[count_col], 0) != len(ids):
                raise WorkbookImportError("maxed_player_count_mismatch")
            banks.append(WorkbookBank(who, None if masters_only else quotas[0], quotas[-1],
                                      tuple(ids), row[count_col], provenance))
            records += 1
    elif sheet == "ChoixSommaire":
        if len(rows) < 4 or width < 2 or not any(v is not None for v in rows[2][1:]):
            raise _UnsupportedLayout
        columns = [j for j, v in enumerate(rows[2]) if j > 0 and v is not None]
        usage_participants = [participant(rows[2][j]) for j in columns]
        if len(set(usage_participants)) != len(usage_participants):
            raise WorkbookImportError("duplicate_usage_participant")
        for row in rows[4:]:
            if not any(v is not None for v in row):
                continue
            player_id = identity(row[0])
            unique(player_id)
            if any(v is not None and j not in columns for j, v in enumerate(row[1:], 1)):
                raise _UnsupportedLayout
            for j in columns:
                count = number(row[j], 5)
                usage.append(WorkbookUsage(participant(rows[2][j]), player_id, count, provenance))
                records += 1
    elif sheet == "Substitutions":
        if width != 5 or rows[0] != ("Participant", "Event", "Withdrawn", "Replacement", "Band"):
            raise _UnsupportedLayout
        for row in rows[1:]:
            if not any(v is not None for v in row):
                continue
            who = participant(row[0])
            if row[1] not in TOURNAMENTS:
                raise WorkbookImportError("invalid_substitution_event")
            old, new = identity(row[2]), identity(row[3])
            if old == new:
                raise WorkbookImportError("invalid_substitution_identity")
            unique((who, row[1], old))
            category = band(row[4])
            replacements.append(WorkbookSubstitution(who, row[1], old, new, category, provenance))
            records += 1
            substitutions += 1
    return (SheetSummary(sheet, records, missing, substitutions, provenance),
            tuple(rankings), tuple(warnings), WorkbookSeason(
                tuple(picks), tuple(standings), tuple(banks), tuple(usage), tuple(replacements),
            ))


def preview_workbook(
    path: str | Path, *, event: Event, year: int, format: str,
    registry: IdentityRegistry, provenance: Provenance,
    limits: WorkbookLimits = WorkbookLimits(),
    season_path: str | Path | None = None,
    season_provenance: Provenance | None = None,
) -> WorkbookPreview:
    """Preview authorized local records without consuming state or logging private rows."""
    try:
        # Parser warnings may quote defined names or other private metadata.
        with warnings.catch_warnings(record=True) as captured:
            warnings.simplefilter("always")
            result = _preview(path, event, year, format, registry, provenance, limits,
                              season_path, season_provenance)
        if captured:
            result = replace(result, warnings=result.warnings + ("workbook_parser_warning",))
        return result
    except WorkbookImportError as error:
        reason = str(error)
    except IdentityResolutionError:
        reason = "identity_requires_review"
    except ImportError:
        reason = "excel_dependency_required"
    except Exception:
        # Third-party parser exceptions can include filenames or private values.
        reason = "invalid_workbook_or_metadata"
    # Raise outside handlers so even __context__ cannot expose private source errors.
    raise WorkbookImportError(reason) from None


def _preview(
    path: str | Path, event: Event, year: int, format: str,
    registry: IdentityRegistry, provenance: Provenance, limits: WorkbookLimits,
    season_path: str | Path | None, season_provenance: Provenance | None,
) -> WorkbookPreview:
    from openpyxl import load_workbook

    event_name = _EVENT_LABELS.get(event.name, event.name)
    expected = "GS128" if event_name in _SLAMS else "M6" if event_name in _SIX_ROUNDS else "M7"
    if (event_name not in TOURNAMENTS or format != expected or type(year) is not int
            or not 1900 <= year <= 2100 or event.start_date.year != year
            or event.end_date.year != year or event.tour is not Tour.ATP):
        raise WorkbookImportError("invalid_supplied_event_year_or_format")
    size, entrants, seed_count, rounds, best_of = _FORMATS[format]
    if event.round_count != rounds or event.best_of != best_of:
        raise WorkbookImportError("event_format_mismatch")
    data = _bounded_bytes(path, limits)
    with ExitStack() as stack:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True,
                           keep_links=False, keep_vba=False)
        stack.callback(wb.close)
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False,
                                 keep_links=False, keep_vba=False)
        stack.callback(formulas.close)
        sources = [(wb, formulas, provenance)]
        if (season_path is None) != (season_provenance is None):
            raise WorkbookImportError("season_path_and_provenance_required_together")
        if season_path is not None and season_provenance is not None:
            season_data = _bounded_bytes(season_path, limits)
            season_wb = load_workbook(BytesIO(season_data), read_only=True, data_only=True,
                                      keep_links=False, keep_vba=False)
            stack.callback(season_wb.close)
            season_formulas = load_workbook(BytesIO(season_data), read_only=True, data_only=False,
                                            keep_links=False, keep_vba=False)
            stack.callback(season_formulas.close)
            sources.append((season_wb, season_formulas, season_provenance))
        if "Draw" not in wb.sheetnames:
            raise WorkbookImportError("missing_Draw")
        rows, _ = _sheet_rows(wb["Draw"])
        _, formula_positions = _sheet_rows(formulas["Draw"])
        if not rows or len(rows[0]) < 4:
            raise WorkbookImportError("unsupported_Draw_layout")
        headers = [i for i, row in enumerate(rows) if row[0] == "Draw"]
        if len(headers) != 1:
            raise WorkbookImportError("unsupported_Draw_layout")
        start = headers[0]
        for label, allowed in zip(rows[start][1:4], (
            (None, "Player", "Joueur", "Nom"),
            (None, "Rank", "Rankings", "Rang", "Classement", "ATP"),
            (None, "Seed", "T\u00eate de s\u00e9rie", "TS"),
        )):
            if label not in allowed:
                raise WorkbookImportError("unsupported_Draw_columns")
        warnings = []
        for i, j in sorted(formula_positions):
            if rows[i][j] is None:
                if "Draw:missing_formula_cache" not in warnings:
                    warnings.append("Draw:missing_formula_cache")
                if i > start and j in (0, 1, 3):
                    raise WorkbookImportError("Draw:required_formula_cache_absent")
        slots = []
        rankings = []
        for row in rows[start + 1:]:
            if row[0] is None:
                if any(v is not None for v in row[1:]):
                    raise WorkbookImportError("draw_row_without_position")
                continue
            position = _integer(row[0])
            name = row[1]
            if name is None or (isinstance(name, str) and name.strip().casefold() == "bye"):
                if row[2] not in (None, 9999, "#N/A") or row[3] is not None:
                    raise WorkbookImportError("invalid_bye_rank_or_seed")
                slots.append(DrawSlot(position=position, kind=EntryKind.BYE, player_id=None))
                continue
            if not isinstance(name, str) or not name.strip():
                raise WorkbookImportError("invalid_draw_player")
            name = name.strip()
            player_id = registry.resolve(source=provenance.source, tour=Tour.ATP, name=name)
            slots.append(DrawSlot(
                position=position, kind=EntryKind.PLAYER, player_id=player_id,
                seed=None if row[3] is None else _integer(row[3]),
            ))
            rankings.append(WorkbookRanking(
                player_id=player_id, sheet="Draw",
                rank=_rank(row[2]),
                provenance=provenance,
            ))
        if format == "M6" and len(slots) == 128:
            if any(slot.kind is not EntryKind.BYE or slot.position != i
                   for i, slot in enumerate(slots[64:], 65)) or any(
                       any(value is not None for value in row[4:])
                       for row in rows[start + 1:] if row[0] is not None and row[0] > 64
                   ):
                raise WorkbookImportError("unsupported_six_round_padding")
            slots = slots[:64]
            warnings.append("trailing_bye_padding_removed")
        if len(slots) != size or sum(s.player_id is not None for s in slots) != entrants:
            raise WorkbookImportError("draw_slot_or_entrant_count_mismatch")
        if sorted(s.seed for s in slots if s.seed is not None) != list(range(1, seed_count + 1)):
            raise WorkbookImportError("draw_seed_count_mismatch")
        summaries = []
        missing_sheets = []
        unsupported = []
        season = WorkbookSeason()
        for sheet in (*_SUPPORT, "Substitutions"):
            matches = [(book, formula_book, evidence) for book, formula_book, evidence in sources
                       if sheet in book.sheetnames]
            if not matches:
                if sheet in _SUPPORT:
                    missing_sheets.append(sheet)
                continue
            if len(matches) > 1:
                raise WorkbookImportError("ambiguous_support_sheet_sources")
            book, formula_book, evidence = matches[0]
            support_rows, _ = _sheet_rows(book[sheet])
            _, support_formulas = _sheet_rows(formula_book[sheet])
            absent = {(i, j) for i, j in support_formulas if support_rows[i][j] is None}
            if absent:
                warnings.append(sheet + ":missing_formula_cache")
            # Missing key caches cannot be mistaken for blank/no-history rows.
            bank_name_column = (8 if support_rows and len(support_rows[0]) > 2
                                and support_rows[0][2] == "Tournois Masters" else 12)
            if any(
                (sheet == "Rankings" and j == 1)
                or (sheet == "Choix" and (j < 2 or j % 2 == 0))
                or (sheet == "Classement" and j < 2)
                or (sheet == "Banques (Solde restant)" and j in (1, bank_name_column))
                or (sheet == "ChoixSommaire" and (j == 0 or i == 2))
                or (sheet == "Substitutions" and j < 4)
                for i, j in absent
            ):
                unsupported.append(sheet)
                continue
            try:
                summary, extra_rankings, extra_warnings, extra_season = _support_summary(
                    sheet, support_rows, registry, evidence,
                )
            except _UnsupportedLayout:
                unsupported.append(sheet)
                continue
            summaries.append(summary)
            rankings.extend(extra_rankings)
            warnings.extend(extra_warnings)
            season = WorkbookSeason(
                season.picks + extra_season.picks, season.standings + extra_season.standings,
                season.banks + extra_season.banks, season.usage + extra_season.usage,
                season.substitutions + extra_season.substitutions,
            )
        return WorkbookPreview(
            draw=Draw(event=event, slots=tuple(slots), provenance=provenance),
            rankings=tuple(rankings), summaries=tuple(summaries),
            missing_sheets=tuple(missing_sheets), unsupported_sheets=tuple(unsupported),
            warnings=tuple(warnings),
            registry_revision=registry.revision, registry_digest=registry.digest,
            season=season,
        )
