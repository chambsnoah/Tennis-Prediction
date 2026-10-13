"""Public preview boundary, using only runtime-generated synthetic workbooks."""

from datetime import date, datetime, timezone
from dataclasses import FrozenInstanceError, replace
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
import pytest

from tennis_api.models.canonical import (
    Event, EventId, Player, PlayerId, Provenance, Surface, Tour,
)
from tennis_api.models.identity import IdentityRegistry, ReviewedAlias
from tennis_pool.workbooks import WorkbookImportError, WorkbookLimits, preview_workbook


NOW = datetime(2026, 1, 1, tzinfo=timezone.utc)
EVIDENCE = Provenance(
    source="synthetic-workbook", source_version="fixture-1", source_record_id=None,
    observed_at=None, available_at=NOW, ingested_at=NOW,
)


@pytest.fixture
def registry():
    players = tuple(
        Player(player_id=PlayerId(f"fixture:p{i}"), display_name=f"Player {i}",
               tour=Tour.ATP, provenance=EVIDENCE)
        for i in range(1, 130)
    )
    return IdentityRegistry(
        revision="synthetic-review-1", players=players,
        aliases=tuple(
            ReviewedAlias(source=EVIDENCE.source, player_id=p.player_id,
                          tour=Tour.ATP, name=p.display_name,
                          reviewed_by="fixture-reviewer", reviewed_at=NOW,
                          provenance=EVIDENCE)
            for p in players
        ),
    )


def event(name="Australie", rounds=7, best_of=5):
    return Event(
        event_id=EventId("fixture:event"), name=name, tour=Tour.ATP,
        start_date=date(2026, 1, 10), end_date=date(2026, 1, 20),
        surface=Surface.HARD, best_of=best_of, round_count=rounds,
        level="grand_slam" if best_of == 5 else "masters", provenance=EVIDENCE,
    )


def workbook(slots=128, entrants=128, seeds=32):
    wb = Workbook()
    ws = wb.active
    ws.title = "Draw"
    ws.append(["Draw", "Player", "Rank", "Seed"])
    player = 0
    byes = slots - entrants
    for position in range(1, slots + 1):
        if position % 2 == 0 and byes:
            ws.append([position, "BYE", 9999, None])
            byes -= 1
        else:
            player += 1
            ws.append([position, f"Player {player}", None if player == 1 else player,
                       player if player <= seeds else None])
    return wb


def preview(tmp_path, registry, wb=None, **kwargs):
    path = tmp_path / "synthetic.xlsx"
    wb = wb or workbook()
    wb.save(path)
    wb.close()
    return preview_workbook(
        path, event=kwargs.pop("event", event()), year=2026,
        format=kwargs.pop("format", "GS128"), registry=registry,
        provenance=EVIDENCE, **kwargs,
    )


def test_slam_preview_returns_canonical_draw_and_unknown_rank_without_mutation(tmp_path, registry):
    before = registry.digest
    result = preview(tmp_path, registry)
    assert len(result.draw.slots) == 128
    assert result.draw.slots[0].player_id == PlayerId("fixture:p1")
    assert result.draw.slots[0].seed == 1
    assert result.rankings[0].rank is None
    assert result.draw.provenance == EVIDENCE
    assert result.registry_digest == before == registry.digest
    assert result.registry_revision == "synthetic-review-1"
    assert "Rankings" in result.missing_sheets
    assert "ChoixSommaire" in result.missing_sheets
    assert result.summaries == ()


@pytest.mark.parametrize("name,format,slots,entrants,seeds,rounds,best_of", [
    (name, "GS128", 128, 128, 32, 7, 5)
    for name in ("Australie", "Roland-Garros", "Wimbledon", "US Open")
] + [
    (name, "M7", 128, 96, 32, 7, 3)
    for name in ("Indian Wells", "Miami", "Madrid", "Rome", "Montr\u00e9al", "Cincinnati", "Shanghai")
] + [(name, "M6", 64, 56, 16, 6, 3) for name in ("Monte Carlo", "Paris")])
def test_all_thirteen_events_have_explicit_validated_formats(
    tmp_path, registry, name, format, slots, entrants, seeds, rounds, best_of,
):
    result = preview(tmp_path, registry, workbook(slots, entrants, seeds),
                     event=event(name, rounds, best_of), format=format)
    assert len(result.draw.slots) == slots
    assert sum(s.player_id is not None for s in result.draw.slots) == entrants
    assert sum(s.seed is not None for s in result.draw.slots) == seeds
    with pytest.raises(FrozenInstanceError):
        result.registry_revision = "mutated"


@pytest.mark.parametrize("metadata", [
    {"event": event("Unknown")}, {"format": "M7"}, {"format": "unknown"},
    {"event": event(rounds=6)}, {"event": event(best_of=3)},
    {"event": replace(event(), tour=Tour.WTA)},
    {"event": replace(event(), start_date=date(2025, 12, 31))},
])
def test_rejects_invalid_explicit_event_metadata(tmp_path, registry, metadata):
    with pytest.raises(ValueError):
        preview(tmp_path, registry, **metadata)


def test_rejects_wrong_entrant_and_seed_counts(tmp_path, registry):
    for entrants, seeds in ((127, 32), (128, 31)):
        with pytest.raises(ValueError):
            preview(tmp_path, registry, workbook(128, entrants, seeds))


def test_only_safe_trailing_six_round_padding_is_removed(tmp_path, registry):
    wb = workbook(64, 56, 16)
    for position in range(65, 129):
        wb["Draw"].append([position, "BYE", 9999, None])
    result = preview(tmp_path, registry, wb, event=event("Paris", 6, 3), format="M6")
    assert len(result.draw.slots) == 64
    assert "trailing_bye_padding_removed" in result.warnings


@pytest.mark.parametrize("cell,value", [
    ("A3", 1), ("A2", 0), ("A2", 1.5), ("A2", True),
    ("B3", "Player 1"), ("B2", "Unreviewed private label"),
    ("C2", 0), ("C2", -1), ("C2", "invalid"), ("C2", True),
    ("C2", 9999),
    ("D3", 1), ("D2", 33), ("D2", 1.5),
    ("A1", "Unsupported"), ("B1", "Seed"), ("D1", "Rank"),
])
def test_rejects_invalid_draw_values_without_echoing_data(tmp_path, registry, cell, value):
    wb = workbook()
    wb["Draw"][cell] = value
    with pytest.raises(WorkbookImportError) as error:
        preview(tmp_path, registry, wb)
    assert "Unreviewed private label" not in str(error.value)
    assert error.value.__cause__ is None


def test_byes_cannot_have_rank_or_seed_and_padding_cannot_contain_players(tmp_path, registry):
    for cell, value in (("C3", 5), ("D3", 1)):
        wb = workbook(64, 56, 16)
        wb["Draw"][cell] = value
        with pytest.raises(WorkbookImportError):
            preview(tmp_path, registry, wb, event=event("Paris", 6, 3), format="M6")
    wb = workbook(64, 56, 16)
    for position in range(65, 129):
        wb["Draw"].append([position, "Player 129" if position == 65 else "BYE",
                           129 if position == 65 else 9999, None])
    with pytest.raises(WorkbookImportError, match="padding"):
        preview(tmp_path, registry, wb, event=event("Paris", 6, 3), format="M6")


def test_formula_rank_stays_unknown_but_missing_identity_cache_fails(tmp_path, registry):
    wb = workbook()
    wb["Draw"]["C2"] = "=1+1"
    result = preview(tmp_path, registry, wb)
    assert result.rankings[0].rank is None
    assert "Draw:missing_formula_cache" in result.warnings
    wb = workbook(128, 96, 32)
    wb["Draw"]["B3"] = '=IF(TRUE,"Player 1","")'
    with pytest.raises(WorkbookImportError, match="formula_cache"):
        preview(tmp_path, registry, wb, event=event("Miami", 7, 3), format="M7")


@pytest.mark.parametrize("limits", [
    WorkbookLimits(file_bytes=10), WorkbookLimits(zip_bytes=100),
    WorkbookLimits(zip_members=2), WorkbookLimits(rows=100),
    WorkbookLimits(cells=100), WorkbookLimits(columns=3),
])
def test_bounds_generated_workbooks_before_parsing(tmp_path, registry, limits):
    with pytest.raises(WorkbookImportError, match="limit"):
        preview(tmp_path, registry, limits=limits)


def support_sheets(wb):
    ws = wb.create_sheet("Rankings")
    ws.append(["Rank", "Player"])
    ws.append([None, "Player 1"])
    ws.append([2, "Player 2"])
    ws = wb.create_sheet("Choix")
    ws.append(["Participants", "Pick"] + [v for name in (
        "Australie", "Indian Wells", "Miami", "Monte Carlo", "Madrid", "Rome",
        "Roland-Garros", "Wimbledon", "Montr\u00e9al", "Cincinnati", "US Open", "Shanghai", "Paris",
    ) for v in (name, "Band")])
    ws.append(["Private Synthetic Participant", "Joueur 1", "Player 1", "#1-4"])
    ws.append([None, "Substitut 1", "Player 2", "#1-4"])
    ws = wb.create_sheet("Classement")
    ws.cell(4, 1, "Position")
    ws.cell(4, 2, "Participants")
    ws.cell(4, 17, "Total")
    ws.append([1, "Private Synthetic Participant"] + [10] * 13 + [None, 130])
    ws = wb.create_sheet("Banques (Solde restant)")
    ws.cell(1, 3, "Tournois Grand Chelem")
    ws.cell(1, 7, "Tournois Masters")
    for offset in (3, 7):
        for i, band in enumerate(("1-4", "5-8", "9-16", "17-32")):
            ws.cell(2, offset + i, band)
    ws.append([None, "Private Synthetic Participant"] + [2] * 8 + [None, 1, "Player 1"])
    ws = wb.create_sheet("ChoixSommaire")
    ws.cell(3, 2, "Private Synthetic Participant")
    ws.cell(5, 1, "Player 1")
    ws.cell(5, 2, 5)
    ws.cell(6, 1, "Player 2")
    ws.cell(6, 2, "=1+1")
    ws = wb.create_sheet("Substitutions")
    ws.append(["Participant", "Event", "Withdrawn", "Replacement", "Band"])
    ws.append(["Private Synthetic Participant", "Australie", "Player 1", "Player 2", "#1-4"])
    return wb


def test_available_support_layouts_return_only_validated_count_summaries(tmp_path, registry):
    result = preview(tmp_path, registry, support_sheets(workbook()))
    summaries = {s.sheet: s for s in result.summaries}
    assert {name: s.records for name, s in summaries.items()} == {
        "Rankings": 2, "Choix": 2, "Classement": 1,
        "Banques (Solde restant)": 1, "ChoixSommaire": 2, "Substitutions": 1,
    }
    assert summaries["Choix"].substitutions == 1
    assert summaries["ChoixSommaire"].missing_values == 1
    assert result.missing_sheets == ()
    assert "ChoixSommaire:missing_formula_cache" in result.warnings
    assert result.unsupported_sheets == ()
    assert "Private Synthetic Participant" not in repr(result)
    assert [(r.player_id, r.rank) for r in result.rankings if r.sheet == "Rankings"] == [
        (PlayerId("fixture:p1"), None), (PlayerId("fixture:p2"), 2),
    ]


def test_separate_season_summary_has_its_own_provenance(tmp_path, registry):
    season = support_sheets(Workbook())
    season_path = tmp_path / "synthetic-season.xlsx"
    season.save(season_path)
    season.close()
    evidence = replace(EVIDENCE, source_record_id="synthetic-season")
    result = preview(tmp_path, registry, season_path=season_path, season_provenance=evidence)
    assert result.missing_sheets == ()
    assert all(s.provenance == evidence for s in result.summaries)


def test_preview_retains_local_season_records_without_consuming_them(tmp_path, registry):
    result = preview(tmp_path, registry, support_sheets(workbook()))
    season = result.season
    assert season.picks[0].participant == "Private Synthetic Participant"
    assert season.picks[0].event == "Australie"
    assert season.picks[0].slot == "Joueur 1"
    assert season.picks[0].player_id == PlayerId("fixture:p1")
    assert season.picks[0].band == "1-4"
    assert not season.picks[0].is_substitute
    assert season.picks[1].is_substitute
    assert season.standings[0].position == 1
    assert season.standings[0].event_points == (10,) * 13
    assert season.standings[0].bonus is None
    assert season.standings[0].total == 130
    assert season.banks[0].slam == (2, 2, 2, 2)
    assert season.banks[0].masters == (2, 2, 2, 2)
    assert season.banks[0].maxed_players == (PlayerId("fixture:p1"),)
    assert season.usage[0].count == 5
    assert season.usage[1].count is None
    assert season.substitutions[0].withdrawn_id == PlayerId("fixture:p1")
    assert season.substitutions[0].replacement_id == PlayerId("fixture:p2")
    assert all(record.provenance == EVIDENCE for record in (
        *season.picks, *season.standings, *season.banks, *season.usage, *season.substitutions,
    ))
    assert "Private Synthetic Participant" not in repr(result)
    with pytest.raises(FrozenInstanceError):
        season.usage[0].count = 0


@pytest.mark.parametrize("headings", [("Paris", "Australie"), ("Unknown", "Indian Wells")])
def test_standings_event_heading_mismatch_cannot_misattribute_points(tmp_path, registry, headings):
    wb = support_sheets(workbook())
    wb["Classement"]["C4"], wb["Classement"]["D4"] = headings
    wb["Classement"]["C5"], wb["Classement"]["D5"] = 111, 222
    result = preview(tmp_path, registry, wb)
    assert "Classement" in result.unsupported_sheets
    assert result.season.standings == ()


@pytest.mark.parametrize("valid", [True, False])
def test_private_parser_warning_is_redacted_on_success_and_failure(tmp_path, registry, valid, capsys, recwarn):
    wb = workbook()
    wb.defined_names.add(DefinedName(
        "_xlnm.Print_Area", attr_text="PRIVATE_SYNTHETIC_SECRET", localSheetId=0,
    ))
    if not valid:
        wb["Draw"]["B2"] = "Unknown"
        with pytest.raises(WorkbookImportError):
            preview(tmp_path, registry, wb)
    else:
        result = preview(tmp_path, registry, wb)
        assert "workbook_parser_warning" in result.warnings
    assert not any("PRIVATE_SYNTHETIC_SECRET" in str(w.message) for w in recwarn)
    assert "PRIVATE_SYNTHETIC_SECRET" not in capsys.readouterr().err


def test_missing_full_season_quota_cache_preserves_other_bank_observations(tmp_path, registry):
    wb = support_sheets(workbook())
    wb["Banques (Solde restant)"]["I3"] = "=1+1"
    result = preview(tmp_path, registry, wb)
    assert "Banques (Solde restant)" not in result.unsupported_sheets
    assert result.season.banks[0].masters == (2, 2, None, 2)
    assert "Banques (Solde restant):missing_formula_cache" in result.warnings


@pytest.mark.parametrize("sheet,cell,value", [
    ("Rankings", "A3", 0), ("Rankings", "B3", "Player 1"),
    ("Choix", "C2", "Unreviewed private label"), ("Choix", "D2", "invalid"),
    ("Banques (Solde restant)", "C3", -1), ("Banques (Solde restant)", "C3", 1.5),
    ("Banques (Solde restant)", "L3", 2),
    ("ChoixSommaire", "B5", 6), ("ChoixSommaire", "B5", -1),
    ("Classement", "A5", 0), ("Classement", "Q5", "Private bad total"),
    ("Substitutions", "D2", "Player 1"), ("Substitutions", "B2", "Unknown"),
])
def test_rejects_invalid_support_data_privately(tmp_path, registry, sheet, cell, value):
    wb = support_sheets(workbook())
    wb[sheet][cell] = value
    with pytest.raises(WorkbookImportError) as error:
        preview(tmp_path, registry, wb)
    assert "Private" not in str(error.value)
    assert "Unreviewed" not in str(error.value)


def test_unsupported_support_layout_is_explicit_not_empty_success(tmp_path, registry):
    wb = workbook()
    wb.create_sheet("ChoixSommaire").append(["Unknown layout", "Private label"])
    result = preview(tmp_path, registry, wb)
    assert result.unsupported_sheets == ("ChoixSommaire",)
    assert "ChoixSommaire" not in result.missing_sheets
    assert result.summaries == ()


def rewrite_archive(wb, transform, extra=()):
    source = BytesIO()
    wb.save(source)
    wb.close()
    output = BytesIO()
    with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in original.infolist():
            target.writestr(member.filename, transform(member.filename, original.read(member)))
        for name, content in extra:
            target.writestr(name, content)
    return output.getvalue()


def preview_bytes(tmp_path, registry, data, suffix="xlsx", **kwargs):
    path = tmp_path / f"synthetic.{suffix}"
    path.write_bytes(data)
    return preview_workbook(path, event=event(), year=2026, format="GS128",
                            registry=registry, provenance=EVIDENCE, **kwargs)


@pytest.mark.parametrize("dimension", [b'<dimension ref="A1:A1"/>', b""])
def test_underreported_or_absent_dimensions_do_not_truncate_data(tmp_path, registry, dimension):
    data = rewrite_archive(support_sheets(workbook()), lambda name, content:
        re.sub(rb'<dimension[^>]*/>', dimension, content)
        if name.startswith("xl/worksheets/") else content)
    result = preview_bytes(tmp_path, registry, data)
    assert len(result.draw.slots) == 128
    assert {s.sheet: s.records for s in result.summaries}["ChoixSommaire"] == 2


@pytest.mark.parametrize("transform,reason", [
    (lambda b: re.sub(rb'<dimension[^>]*/>', b'<dimension ref="A1:XFD1048576"/>', b), "limit"),
    (lambda b: b'<!DOCTYPE worksheet [<!ENTITY secret "PRIVATE">]>' + b, "unsafe_xml"),
    (lambda b: b"<worksheet><sheetData><row", "invalid_workbook"),
    (lambda b: b.replace(b'r="A2"', b'r="A3"'), "invalid"),
])
def test_rejects_adversarial_sheet_xml(tmp_path, registry, transform, reason):
    data = rewrite_archive(workbook(), lambda name, content:
        transform(content) if name == "xl/worksheets/sheet1.xml" else content)
    with pytest.raises(WorkbookImportError, match=reason):
        preview_bytes(tmp_path, registry, data)


def test_malformed_files_and_duplicate_zip_members_are_redacted(tmp_path, registry):
    for data in (b"Private confidential content", b"PK\x03\x04corrupt"):
        with pytest.raises(WorkbookImportError) as error:
            preview_bytes(tmp_path, registry, data)
        assert "Private" not in str(error.value)
        assert error.value.__cause__ is None
    with pytest.warns(UserWarning, match="Duplicate"):
        data = rewrite_archive(workbook(), lambda name, content: content,
                               extra=(("xl/workbook.xml", b"Private content"),))
    with pytest.raises(WorkbookImportError, match="duplicate_zip"):
        preview_bytes(tmp_path, registry, data)


def test_xlsm_macros_are_opaque_data_and_external_relationships_are_not_fetched(tmp_path, registry):
    def macro_and_link_parts(name, content):
        if name == "[Content_Types].xml":
            return content.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            )
        if name == "xl/workbook.xml":
            return content.replace(b"</workbook>",
                b'<externalReferences><externalReference '
                b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                b'r:id="rId99"/></externalReferences></workbook>')
        if name == "xl/_rels/workbook.xml.rels":
            return content.replace(b"</Relationships>",
                b'<Relationship Id="rId99" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" '
                b'Target="/xl/externalLinks/externalLink1.xml"/>'
                b'<Relationship Id="rId100" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
                b'Target="vbaProject.bin"/></Relationships>')
        return content

    data = rewrite_archive(workbook(), macro_and_link_parts,
        extra=(("xl/vbaProject.bin", b"synthetic, never executed"),
               ("xl/externalLinks/externalLink1.xml",
                b'<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                b'<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                b'r:id="rId1"><sheetNames><sheetName val="PRIVATE"/></sheetNames>'
                b'<sheetDataSet/></externalBook></externalLink>'),
               ("xl/externalLinks/_rels/externalLink1.xml.rels",
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId1" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
                b'Target="https://invalid.example/PRIVATE" '
                b'TargetMode="External"/></Relationships>')))
    result = preview_bytes(tmp_path, registry, data, suffix="xlsm")
    assert len(result.draw.slots) == 128
    assert "PRIVATE" not in repr(result)


def test_existing_numeric_formula_cache_is_used_not_recalculated(tmp_path, registry):
    wb = workbook()
    wb["Draw"]["C2"] = "=1+1"
    data = rewrite_archive(wb, lambda name, content:
        re.sub(rb"<f>1\+1</f><v(?:></v>|\s*/>)", b"<f>1+1</f><v>7</v>", content)
        if name == "xl/worksheets/sheet1.xml" else content)
    result = preview_bytes(tmp_path, registry, data)
    assert result.rankings[0].rank == 7
    assert "Draw:missing_formula_cache" not in result.warnings


@pytest.mark.parametrize("missing_names_cache", [False, True])
def test_masters_only_banks_report_absent_slam_quotas(tmp_path, registry, missing_names_cache):
    wb = workbook()
    ws = wb.create_sheet("Banques (Solde restant)")
    ws.cell(1, 3, "Tournois Masters")
    for i, band in enumerate(("1-4", "5-8", "9-16", "17-32"), 3):
        ws.cell(2, i, band)
    ws.append([None, "Private Synthetic Participant", 9, 9, 9, 7, None, 0, None])
    if missing_names_cache:
        ws["I3"] = '=IF(TRUE,"Player 1","")'
    result = preview(tmp_path, registry, wb)
    if missing_names_cache:
        assert "Banques (Solde restant)" in result.unsupported_sheets
        assert result.season.banks == ()
        return
    assert result.summaries[0].records == 1
    assert result.season.banks[0].slam is None
    assert result.season.banks[0].masters == (9, 9, 9, 7)
    assert "Banques (Solde restant):slam_quotas_absent" in result.warnings


def test_support_identity_formula_cache_absence_cannot_erase_history(tmp_path, registry):
    wb = support_sheets(workbook())
    wb["Choix"]["C2"] = '=IF(TRUE,"Player 1","")'
    wb["Choix"]["D2"] = None
    result = preview(tmp_path, registry, wb)
    assert "Choix" in result.unsupported_sheets
    assert not any(s.sheet == "Choix" for s in result.summaries)


@pytest.mark.parametrize("year", [None, True, 2025, 0, 2026.0, 2101])
def test_invalid_year_is_rejected_at_public_boundary(tmp_path, registry, year):
    wb = workbook()
    path = tmp_path / "synthetic.xlsx"
    wb.save(path)
    wb.close()
    with pytest.raises(WorkbookImportError, match="year"):
        preview_workbook(path, event=event(), year=year, format="GS128",
                         registry=registry, provenance=EVIDENCE)


@pytest.mark.parametrize("name", ["Australian Open", "French Open", "Canadian Masters", "Toronto", "Montreal"])
def test_explicit_public_event_labels_map_to_one_reference_event(tmp_path, registry, name):
    slam = name in ("Australian Open", "French Open")
    result = preview(tmp_path, registry, workbook(128, 128 if slam else 96, 32),
                     event=event(name, 7, 5 if slam else 3), format="GS128" if slam else "M7")
    assert result.draw.event.name == name


def test_draw_rows_without_positions_and_populated_padding_are_not_silently_discarded(tmp_path, registry):
    wb = workbook()
    wb["Draw"].append([None, "Player 129", 129, None])
    with pytest.raises(WorkbookImportError, match="position"):
        preview(tmp_path, registry, wb)
    wb = workbook(64, 56, 16)
    for position in range(65, 129):
        wb["Draw"].append([position, "BYE", 9999, None, "Player 1"])
    with pytest.raises(WorkbookImportError, match="padding"):
        preview(tmp_path, registry, wb, event=event("Paris", 6, 3), format="M6")


def test_hidden_support_sheets_are_still_bounded(tmp_path, registry):
    wb = workbook()
    ws = wb.create_sheet("Private hidden sheet")
    ws.sheet_state = "hidden"
    ws.cell(10_001, 1, "Private content")
    with pytest.raises(WorkbookImportError, match="limit"):
        preview(tmp_path, registry, wb)


@pytest.mark.parametrize("cell,value", [
    ("C3", 5), ("J3", 8), ("M3", "Player 1,Player 1"), ("M3", "Unreviewed private label"),
])
def test_bank_capacities_and_maxed_identities_are_validated(tmp_path, registry, cell, value):
    wb = support_sheets(workbook())
    wb["Banques (Solde restant)"][cell] = value
    with pytest.raises(WorkbookImportError):
        preview(tmp_path, registry, wb)


def test_duplicate_usage_players_and_participant_blocks_are_rejected(tmp_path, registry):
    for sheet, cell, value in (("ChoixSommaire", "A6", "Player 1"),
                               ("Choix", "A3", "Private Synthetic Participant")):
        wb = support_sheets(workbook())
        wb[sheet][cell] = value
        with pytest.raises(WorkbookImportError, match="duplicate"):
            preview(tmp_path, registry, wb)


def test_season_provenance_is_required_and_ambiguous_sources_fail(tmp_path, registry):
    season_path = tmp_path / "synthetic-season.xlsx"
    wb = support_sheets(Workbook())
    wb.save(season_path)
    wb.close()
    with pytest.raises(WorkbookImportError, match="provenance"):
        preview(tmp_path, registry, season_path=season_path)
    with pytest.raises(WorkbookImportError, match="ambiguous"):
        preview(tmp_path, registry, support_sheets(workbook()),
                season_path=season_path, season_provenance=EVIDENCE)


def test_optional_excel_dependency_has_actionable_redacted_error(tmp_path, registry, monkeypatch):
    import builtins

    original_import = builtins.__import__

    def without_excel(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("Private dependency error")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_excel)
    with pytest.raises(WorkbookImportError, match="excel_dependency_required"):
        preview(tmp_path, registry)


@pytest.mark.parametrize("failure", [None, "validation", "second_load"])
def test_workbook_resources_close_on_success_and_failure(tmp_path, registry, monkeypatch, failure):
    # Observe external workbook resource lifetimes, not adapter implementation state.
    import openpyxl

    original_load = openpyxl.load_workbook
    closed = []
    opened = []

    def tracked_load(*args, **kwargs):
        if failure == "second_load" and opened:
            raise ValueError("Private parser detail")
        book = original_load(*args, **kwargs)
        opened.append(book)
        original_close = book.close

        def close():
            closed.append(book)
            original_close()

        book.close = close
        return book

    monkeypatch.setattr(openpyxl, "load_workbook", tracked_load)
    wb = workbook()
    if failure == "validation":
        wb["Draw"]["D2"] = 33
    if failure is None:
        preview(tmp_path, registry, wb)
    else:
        with pytest.raises(WorkbookImportError) as error:
            preview(tmp_path, registry, wb)
        assert "Private" not in str(error.value)
        assert error.value.__context__ is None
    assert opened and len(opened) == len(closed)
    assert {id(book) for book in opened} == {id(book) for book in closed}


def test_non_regular_input_is_rejected_without_reading(tmp_path, registry):
    import os

    path = tmp_path / "synthetic.xlsx"
    os.mkfifo(path)
    with pytest.raises(WorkbookImportError, match="regular_file"):
        preview_workbook(path, event=event(), year=2026, format="GS128",
                         registry=registry, provenance=EVIDENCE)


def test_relocated_worksheet_cannot_bypass_preflight_bounds(tmp_path, registry):
    # Relationship targets, not the ZIP filename alone, choose loaded worksheets.
    source = BytesIO()
    wb = workbook()
    wb.save(source)
    wb.close()
    with ZipFile(source) as original:
        worksheet = original.read("xl/worksheets/sheet1.xml")
    data = rewrite_archive(workbook(), lambda name, content:
        content.replace(b"/xl/worksheets/sheet1.xml", b"/xl/private-data.bin")
        if name == "xl/_rels/workbook.xml.rels" else content,
        extra=(("xl/private-data.bin", worksheet),))
    with pytest.raises(WorkbookImportError, match="worksheet_location"):
        preview_bytes(tmp_path, registry, data)
