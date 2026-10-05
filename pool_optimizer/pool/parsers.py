"""Readers for the pool's Excel workbooks.

Two file families:
- Tournament workbooks ("10 - Cincinnati.xlsm", "11 - US Open.xlsm", ...):
  sheets Draw, Rankings, Classement, Choix, <TournamentName>.
- Season summary ("ChoixTournois_2026.xlsx"): sheets Banques (Solde restant),
  ChoixSommaire, Tableaux.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field

from openpyxl import load_workbook

GS_BANDS = ("1-4", "5-8", "9-16", "17-32")
BYE_RANK = 9999

TOURNAMENTS = [
    "Australie", "Indian Wells", "Miami", "Monte Carlo", "Madrid", "Rome",
    "Roland-Garros", "Wimbledon", "Montréal", "Cincinnati", "US Open",
    "Shanghai", "Paris",
]
GRAND_SLAMS = {"Australie", "Roland-Garros", "Wimbledon", "US Open"}
MASTERS_6R = {"Monte Carlo", "Paris"}  # 16 seeds, 6 rounds


def norm(name: str) -> str:
    """Accent/case/space-insensitive key for joining names across sheets."""
    if name is None:
        return ""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


@dataclass
class DrawEntry:
    position: int          # 1-based slot in the bracket
    name: str
    atp_rank: int          # world rank at tournament start (9999 for BYE)
    seed: int | None       # 1..32 for seeds, None otherwise

    @property
    def is_bye(self) -> bool:
        return self.atp_rank == BYE_RANK or norm(self.name) == "bye"


@dataclass
class Draw:
    entries: list[DrawEntry]                      # ordered by position
    # actual results, if the tournament has been played (for back-testing):
    # match -> winner, keyed by (round_no, frozenset({name_a, name_b}))
    results: dict[tuple[int, frozenset], str] = field(default_factory=dict)

    @property
    def n_rounds(self) -> int:
        # played tournaments: trust the recorded rounds (sheets for 6-round
        # events pad the slot list with extra BYEs, so slot count lies)
        if self.results:
            return max(rnd for rnd, _ in self.results)
        # unplayed: derive from the number of real entrants (56 players ->
        # 64-bracket -> 6 rounds), never from the padded slot count
        players = sum(1 for e in self.entries if not e.is_bye)
        n, size = 0, 1
        while size < players:
            size *= 2
            n += 1
        return n


def load_draw(path: str) -> Draw:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Draw"]
    rows = list(ws.iter_rows(values_only=True))

    # locate the header row ("Draw" in column A) and read entries below it
    start = None
    for i, r in enumerate(rows):
        if r and str(r[0]).strip() == "Draw":
            start = i + 1
            break
    if start is None:
        raise ValueError(f"No draw header found in {path}")

    entries: list[DrawEntry] = []
    for r in rows[start:]:
        if not r or r[0] is None:
            break
        try:
            pos = int(r[0])
        except (TypeError, ValueError):
            break
        name = str(r[1]).strip() if r[1] else "BYE"
        try:
            rank = int(r[2])
        except (TypeError, ValueError):  # missing or '#N/A'
            # unknown rank for a real entrant (protected ranking etc.):
            # a mid-tier guess distorts 'better player' bonus calls less
            # than pretending he's the world's worst
            rank = BYE_RANK if norm(name) == "bye" else 200
        try:
            seed = int(r[3])
        except (TypeError, ValueError):
            seed = None
        entries.append(DrawEntry(pos, name, rank, seed))

    n = len(entries)
    if n & (n - 1):
        raise ValueError(f"Draw size {n} is not a power of two in {path}")

    draw = Draw(entries)

    # results: each entrant row lists (opponent, match#, winner) triplets per
    # round; after the entrant is eliminated the chain continues with the
    # matches of whoever advanced on this bracket line, so track that winner
    for i, e in enumerate(entries):
        r = rows[start + i]
        current = norm(e.name)
        rnd = 0
        c = 4
        while c < len(r) - 1:
            opp = r[c]
            rnd += 1
            if opp is None:
                break
            # rounds are (opponent, match#, winner) triplets, except the final
            # which is (opponent, winner) with no match number
            mid = r[c + 1]
            is_final_pair = not isinstance(mid, (int, float))
            winner = mid if is_final_pair else (r[c + 2] if c + 2 < len(r) else None)
            if winner is None:
                break
            w = str(winner).strip()
            if w in ("#VALUE!", "#REF!", ""):
                break
            pair = frozenset({current, norm(str(opp))})
            draw.results[(rnd, pair)] = norm(w)
            current = norm(w)
            if is_final_pair:
                break
            c += 3
    return draw


def load_rankings(path: str) -> dict[str, int]:
    """ATP world ranking at tournament start, keyed by normalized last name.

    Surnames collide (siblings: Tsitsipas, Cerúndolo...); setdefault keeps the
    best-ranked one, which is the player draws refer to by bare surname. Draw
    scoring never uses this (the Draw sheet carries its own rank column)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Rankings"]
    out: dict[str, int] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None or r[1] is None:
            continue
        try:
            rank = int(r[0])
        except (TypeError, ValueError):
            continue
        out.setdefault(norm(r[1]), rank)
    return out


@dataclass
class ParticipantBank:
    name: str
    gs_quota: dict[str, int]        # band -> remaining picks (Slams)
    masters_quota: dict[str, int]   # band -> remaining picks (Masters)
    maxed_players: list[str]        # players already used 5 times


def load_banques(path: str) -> dict[str, ParticipantBank]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Banques (Solde restant)"]
    rows = list(ws.iter_rows(values_only=True))
    # per-event entry forms (ChoixShanghai.xlsx, ...) carry ONLY the bank of
    # that event's pool: "Tournois Masters" in C1, bands in C:F, the maxed
    # count in H and names in I. The season summary has GS C:F + Masters G:J.
    header = str(rows[0][2]).strip() if rows and len(rows[0]) > 2 else ""
    masters_only = header == "Tournois Masters"
    banks: dict[str, ParticipantBank] = {}
    for r in rows[2:]:
        if not r or r[1] is None:
            continue
        name = str(r[1]).strip()
        maxed = []
        if masters_only:
            gs = {b: 0 for b in GS_BANDS}  # Slam pool is over
            ms = {b: int(r[2 + i] or 0) for i, b in enumerate(GS_BANDS)}
            if len(r) > 8 and r[8]:
                maxed = [p.strip() for p in str(r[8]).split(",") if p.strip()]
        else:
            gs = {b: int(r[2 + i] or 0) for i, b in enumerate(GS_BANDS)}
            ms = {b: int(r[6 + i] or 0) for i, b in enumerate(GS_BANDS)}
            if len(r) > 12 and r[12]:  # r[11] is the count; names in r[12]
                maxed = [p.strip() for p in str(r[12]).split(",") if p.strip()]
        banks[name] = ParticipantBank(name, gs, ms, maxed)
    return banks


def load_player_usage(path: str) -> dict[str, dict[str, int]]:
    """Times each participant has used each player (max 5), from ChoixSommaire."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["ChoixSommaire"]
    rows = list(ws.iter_rows(values_only=True))
    participants = {j: str(v).strip() for j, v in enumerate(rows[2]) if v}
    usage: dict[str, dict[str, int]] = {p: {} for p in participants.values()}
    for r in rows[4:]:
        if not r or r[0] is None:
            continue
        player = str(r[0]).strip()
        for j, pname in participants.items():
            v = r[j] if j < len(r) else None
            if v:
                usage[pname][player] = int(v)
    return usage


def load_pick_history(path: str) -> dict[str, dict[str, list[tuple[str, str]]]]:
    """participant -> tournament -> [(player, band)] from a tournament
    workbook's Choix sheet (bands: '#1-4', '#5-8', '#9-16', '#17-32', 'NF')."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Choix"]
    rows = list(ws.iter_rows(values_only=True))
    history: dict[str, dict[str, list[tuple[str, str]]]] = {}
    current = None
    for r in rows:
        if not r:
            continue
        if r[0] and str(r[0]).strip() not in ("Participants", ""):
            current = str(r[0]).strip()
            history.setdefault(current, {t: [] for t in TOURNAMENTS})
        if current and r[1] and str(r[1]).startswith("Joueur"):
            for t_idx, t in enumerate(TOURNAMENTS):
                j = 2 + t_idx * 2
                if j + 1 < len(r) and r[j] and not str(r[j]).startswith("SI("):
                    band = str(r[j + 1]).strip() if r[j + 1] else "NF"
                    history[current][t].append((str(r[j]).strip(), band))
    return history


def load_standings(path: str) -> dict[str, dict]:
    """Season standings from a tournament workbook's Classement sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Classement"]
    rows = list(ws.iter_rows(values_only=True))
    out: dict[str, dict] = {}
    for r in rows[4:]:
        if not r or r[1] is None:
            continue
        try:
            pos = int(r[0])
        except (TypeError, ValueError):
            break
        per_t = {t: r[2 + i] for i, t in enumerate(TOURNAMENTS)}
        out[str(r[1]).strip()] = {
            "rank": pos,
            "per_tournament": per_t,
            "total": r[16],
        }
    return out


def load_pool_results(path: str, tournament: str) -> dict[str, dict]:
    """Ground truth from a tournament workbook's results sheet:
    participant -> {player -> (atp_pts, bonus_pts, total)}. Used for tests."""
    wb = load_workbook(path, read_only=True, data_only=True)
    # sheet names sometimes drop spaces/hyphens ("Indian Wells" -> "IndianWells")
    compact = norm(tournament).replace(" ", "").replace("-", "")
    sheet = next(
        s for s in wb.sheetnames
        if norm(s).replace(" ", "").replace("-", "") == compact
    )
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    out: dict[str, dict] = {}
    current = None
    for r in rows[4:]:
        if not r or r[0] is None:
            continue
        label = str(r[0]).strip()
        base = label.split("_")[0]
        if r[2] is None:
            continue
        current = base
        out.setdefault(current, {})
        player = str(r[2]).strip()
        if not player or player.replace(".", "").isdigit():
            continue  # stray numeric rows in some sheets
        try:
            atp, bonus, total = int(r[4] or 0), int(r[5] or 0), int(r[6] or 0)
        except (TypeError, ValueError):
            continue
        out[current][player] = (atp, bonus, total)
    return out
