"""Discoverable domain coverage using only synthetic, local inputs."""

import json
import random
import subprocess
import sys

import pytest
from openpyxl import Workbook

from pool_optimizer.pool.parsers import (
    GS_BANDS,
    load_banques,
    load_draw,
    load_player_usage,
    load_pool_results,
    load_rankings,
)
from pool_optimizer.pool.results import match_log, score_tournament
from pool_optimizer.pool.scoring import atp_points, band, bonus_points, score_player
from tennis_preds.tennis import Player, PlayerSimple, TennisMatch
from web_interface import server


@pytest.fixture
def preserve_random_state():
    # The core simulator seeds Python's shared RNG, not a per-match instance.
    state = random.getstate()
    yield
    random.setstate(state)


@pytest.mark.parametrize("simple", [False, True], ids=["detailed", "simple"])
def test_core_simulation_is_repeatable_and_accounts_for_points(simple, preserve_random_state):
    if simple:
        players = (PlayerSimple("Synthetic Alpha", 0.68), PlayerSimple("Synthetic Beta", 0.61))
    else:
        players = (
            Player("Synthetic Alpha", 0.65, 0.92, 0.74, 0.55),
            Player("Synthetic Beta", 0.60, 0.94, 0.70, 0.52),
        )

    snapshots = []
    for _ in range(2):
        match = TennisMatch(*players, sets_to_win=2, seed=42)
        match.simulate_match(verbose=False)
        snapshots.append(vars(match).copy())

    assert snapshots[0] == snapshots[1]
    assert max(match.player1_sets_won, match.player2_sets_won) == 2
    assert min(match.player1_sets_won, match.player2_sets_won) < 2
    assert len(match.player1_games_won_per_set) == match.player1_sets_won + match.player2_sets_won
    for games in zip(match.player1_games_won_per_set, match.player2_games_won_per_set):
        winner, loser = max(games), min(games)
        assert (winner == 6 and loser <= 4) or (winner == 7 and loser in (5, 6))
    winner = players[0] if match.player1_sets_won == 2 else players[1]
    assert match.get_match_winner() is winner
    for number in (1, 2):
        total = getattr(match, f"player{number}_total_points_won")
        service = getattr(match, f"player{number}_service_points_won")
        returns = getattr(match, f"player{number}_return_points_won")
        assert total == service + returns
        assert total > 0
    assert match.get_player1_points_won() + match.get_player2_points_won() == pytest.approx(100)
    if not simple:
        points_played = match.player1_total_points_won + match.player2_total_points_won
        assert match.player1_first_serves_played + match.player2_first_serves_played == points_played
        for number in (1, 2):
            assert getattr(match, f"player{number}_second_serves_played") == (
                getattr(match, f"player{number}_first_serves_played")
                - getattr(match, f"player{number}_first_serves_in")
            )


@pytest.mark.parametrize("player1_to_serve", [False, True])
def test_core_one_sided_match_has_exact_score(player1_to_serve, preserve_random_state):
    alpha = PlayerSimple("Synthetic Alpha", 1.0)
    beta = PlayerSimple("Synthetic Beta", 0.0)
    match = TennisMatch(alpha, beta, sets_to_win=2, player1_to_serve=player1_to_serve, seed=7)
    match.simulate_match(verbose=False)

    assert match.get_match_winner() is alpha
    assert match.player1_games_won_per_set == [6, 6]
    assert match.player2_games_won_per_set == [0, 0]
    assert match.player1_total_points_won == 48
    assert match.player2_total_points_won == 0
    assert match.player1_service_points_won == match.player1_return_points_won == 24
    assert match.player1_break_points_converted == 6
    assert match.player1_break_point_chances == 6
    assert match.total_tiebreaks_played == 0


@pytest.mark.parametrize(
    "kind,first_round,final_round,runner_up,champion",
    [("gs", 1, 7, 140, 200), ("m7", 1, 7, 70, 100), ("m6", 1, 6, 70, 100)],
)
def test_pool_atp_points_cover_first_round_final_and_champion(
    kind, first_round, final_round, runner_up, champion
):
    assert atp_points(1, kind) == first_round
    assert atp_points(final_round, kind) == runner_up
    assert atp_points(None, kind) == champion


@pytest.mark.parametrize(
    "seed,expected", [(1, 0), (4, 0), (5, 1), (8, 1), (9, 2), (16, 2), (17, 3), (32, 3), (None, 3)]
)
def test_pool_seed_band_boundaries(seed, expected):
    assert band(seed) == expected


@pytest.mark.parametrize(
    "winner,loser,expected",
    [
        ((None, 80), (1, 1), 50),
        ((8, 12), (2, 2), 30),
        ((1, 1), (8, 12), 0),
        ((4, 4), (1, 1), 20),
        ((1, 1), (4, 4), 0),
        ((None, 80), (None, 40), 10),
        ((None, 40), (None, 80), 0),
        ((None, 5), (17, 100), 10),
        ((17, 100), (None, 5), 0),
    ],
)
def test_pool_upset_bonus_respects_seed_and_rank_order(winner, loser, expected):
    assert bonus_points(*winner, *loser, is_gs=True) == expected
    assert bonus_points(*winner, *loser, is_gs=False) == expected // 2


def test_pool_player_score_accumulates_bonuses():
    result = score_player([(1, 1), (8, 12), (None, 40)], 4, "gs", None, 80)
    assert (result.atp, result.bonus, result.total) == (30, 90, 120)


def test_synthetic_workbook_draw_reconstructs_and_scores_results(tmp_path):
    workbook = Workbook()
    draw_sheet = workbook.active
    draw_sheet.title = "Draw"
    draw_sheet.append(["Synthetic tournament"])
    draw_sheet.append(["Draw", "Player", "Rank", "Seed"])
    draw_sheet.append([1, " Alpha ", 1, 1, "Bravo", 1, "Bravo", "Charlie", "Charlie"])
    draw_sheet.append([2, "Bravo", 12, 8, "Alpha", 1, "Bravo", "Charlie", "Charlie"])
    draw_sheet.append([3, "Charlie", "#N/A", None, "BYE", 2, "Charlie", "Bravo", "Charlie"])
    draw_sheet.append([4, None, None, None, "Charlie", 2, "Charlie", "Bravo", "Charlie"])
    rankings = workbook.create_sheet("Rankings")
    rankings.append(["Rank", "Surname"])
    rankings.append([1, " Alpha "])
    rankings.append([12, "Bravo"])
    rankings.append([99, "BRAVO"])
    rankings.append(["invalid", "Ignored"])
    truth = workbook.create_sheet("RolandGarros")
    truth.cell(4, 7, "Total")
    truth.append(["Participant_1", None, "Alpha", None, 1, 0, 1])
    truth.append(["Participant_2", None, "Bravo", None, 7, 30, 37])
    truth.append(["Participant_3", None, "Charlie", None, 200, 30, 230])
    path = tmp_path / "synthetic_tournament.xlsx"
    workbook.save(path)
    workbook.close()

    draw = load_draw(str(path))
    assert [(entry.position, entry.name, entry.atp_rank, entry.seed) for entry in draw.entries] == [
        (1, "Alpha", 1, 1), (2, "Bravo", 12, 8), (3, "Charlie", 200, None), (4, "BYE", 9999, None)
    ]
    assert draw.entries[-1].is_bye
    assert draw.n_rounds == 2
    assert match_log(draw) == [(1, "bravo", "alpha"), (2, "charlie", "bravo")]
    outcomes = score_tournament(draw, "gs")
    assert set(outcomes) == {"alpha", "bravo", "charlie"}
    assert outcomes["charlie"].lost_in_round is None
    assert outcomes["charlie"].wins == [(8, 12)]
    assert outcomes["bravo"].beaten_by == "charlie"
    assert load_rankings(str(path)) == {"alpha": 1, "bravo": 12}
    recorded = load_pool_results(str(path), "Roland-Garros")["Participant"]
    assert recorded == {"Alpha": (1, 0, 1), "Bravo": (7, 30, 37), "Charlie": (200, 30, 230)}
    for name, expected in recorded.items():
        result = outcomes[name.lower()].result
        assert (result.atp, result.bonus, result.total) == expected


@pytest.mark.parametrize("missing_header", [False, True], ids=["invalid-size", "missing-header"])
def test_synthetic_workbook_rejects_invalid_draw(tmp_path, missing_header):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Draw"
    sheet.append(["Not a draw" if missing_header else "Draw", "Player", "Rank", "Seed"])
    for position in range(1, 4):
        sheet.append([position, f"Synthetic {position}", position, position])
    path = tmp_path / "invalid_draw.xlsx"
    workbook.save(path)
    workbook.close()

    message = "No draw header" if missing_header else "not a power of two"
    with pytest.raises(ValueError, match=message):
        load_draw(str(path))


@pytest.mark.parametrize("masters_only", [False, True], ids=["season-summary", "masters-entry"])
def test_synthetic_workbook_ingests_banks_and_player_usage(tmp_path, masters_only):
    workbook = Workbook()
    banks = workbook.active
    banks.title = "Banques (Solde restant)"
    banks.cell(1, 3, "Tournois Masters" if masters_only else "Tournois Grand Chelem")
    banks.cell(2, 2, "Participant")
    if masters_only:
        banks.append([1, " Synthetic Participant ", 4, 3, 2, 1, None, 2, " Alpha, Bravo "])
    else:
        banks.append([1, " Synthetic Participant ", 8, 7, 6, 5, 4, 3, 2, 1, None, 2, " Alpha, Bravo "])
    usage = workbook.create_sheet("ChoixSommaire")
    usage.cell(3, 2, " Synthetic Participant ")
    usage.cell(4, 1, "Player")
    usage.append(["Alpha", 5])
    usage.append(["Bravo", 2])
    usage.append(["Unused", 0])
    path = tmp_path / "synthetic_summary.xlsx"
    workbook.save(path)
    workbook.close()

    bank = load_banques(str(path))["Synthetic Participant"]
    assert bank.name == "Synthetic Participant"
    assert bank.gs_quota == dict(zip(GS_BANDS, [0, 0, 0, 0] if masters_only else [8, 7, 6, 5]))
    assert bank.masters_quota == dict(zip(GS_BANDS, [4, 3, 2, 1]))
    assert bank.maxed_players == ["Alpha", "Bravo"]
    assert load_player_usage(str(path)) == {"Synthetic Participant": {"Alpha": 5, "Bravo": 2}}


@pytest.fixture
def synthetic_web_client(tmp_path, monkeypatch):
    tournament = tmp_path / "2024" / "wimby2024"
    tournament.mkdir(parents=True)
    (tournament / "players_male.json").write_text(
        json.dumps({"Synthetic Beta": {"seed": 2, "cost": 250}, "Synthetic Alpha": {"seed": 1}}),
        encoding="utf-8",
    )
    (tournament / "bracket_male.txt").write_text(
        "Singles R1 Upcoming Court Synthetic Alpha vs Synthetic Beta\n", encoding="utf-8"
    )
    (tmp_path / "2024" / "empty").mkdir()
    monkeypatch.setattr(server, "PROJECT_ROOT", str(tmp_path))
    monkeypatch.setitem(server.app.config, "TESTING", True)
    with server.app.test_client() as client:
        yield client, tournament


def test_web_lists_synthetic_tournaments_players_and_bracket(synthetic_web_client):
    client, _ = synthetic_web_client
    response = client.get("/api/tournaments")
    assert response.status_code == 200
    assert response.get_json() == [
        {"id": "2024/wimby2024", "name": "Wimbledon 2024", "has_male": True, "has_female": False}
    ]
    response = client.get("/api/players/2024/wimby2024/male")
    assert response.status_code == 200
    assert response.get_json() == [
        {"name": "Synthetic Alpha", "seed": 1, "cost": 0, "p_factor": 0, "n_factor": 0},
        {"name": "Synthetic Beta", "seed": 2, "cost": 250, "p_factor": 0, "n_factor": 0},
    ]
    response = client.get("/api/bracket/2024/wimby2024/male")
    assert response.status_code == 200
    assert response.get_json() == {
        "bracket_text": "Singles R1 Upcoming Court Synthetic Alpha vs Synthetic Beta\n",
        "matches": [{"round": "R1", "status": "Upcoming", "player1": "Synthetic Alpha", "player2": "Synthetic Beta"}],
    }


@pytest.mark.parametrize("resource", ["players", "bracket"])
def test_web_missing_synthetic_data_returns_404(synthetic_web_client, resource):
    client, _ = synthetic_web_client
    response = client.get(f"/api/{resource}/2024/wimby2024/female")
    assert response.status_code == 404
    assert response.get_json() == {"error": f"{resource.title()} file not found"}


def test_web_malformed_synthetic_player_json_returns_error(synthetic_web_client):
    client, tournament = synthetic_web_client
    (tournament / "players_male.json").write_text("{invalid", encoding="utf-8")
    response = client.get("/api/players/2024/wimby2024/male")
    assert response.status_code == 500
    assert "error" in response.get_json()


@pytest.mark.parametrize(
    "endpoint,script,extra_args,timeout",
    [
        ("simulate-match", "match_simulator.py", ["Synthetic Alpha", "Synthetic Beta", "--sets", "3", "--simulations", "500", "--surface", "grass"], 60),
        ("optimize-team", "team_optimizer.py", ["--budget", "100000", "--team-size", "8"], 120),
    ],
)
def test_web_subprocess_success_uses_current_python_and_defaults(
    synthetic_web_client, monkeypatch, endpoint, script, extra_args, timeout
):
    client, _ = synthetic_web_client
    payload = {"tournament_path": "2024/wimby2024", "gender": "male", "player1_name": "Synthetic Alpha", "player2_name": "Synthetic Beta"}
    result = {"synthetic_result": ["Synthetic Alpha"]}
    calls = []

    def fake_run(cmd, **kwargs):
        calls.append((cmd, kwargs))
        return subprocess.CompletedProcess(cmd, 0, stdout=json.dumps(result), stderr="")

    monkeypatch.setattr(server.subprocess, "run", fake_run)
    response = client.post(f"/api/{endpoint}", json=payload)
    assert response.status_code == 200
    assert response.get_json() == result
    assert calls == [
        ([sys.executable, str(server.WEB_ROOT + "/" + script), "2024/wimby2024", "male", *extra_args],
         {"capture_output": True, "text": True, "timeout": timeout})
    ]


@pytest.mark.parametrize("endpoint,label", [("simulate-match", "Simulation"), ("optimize-team", "Optimization")])
@pytest.mark.parametrize("failure", ["stderr", "empty-stderr", "timeout", "invalid-json"])
def test_web_subprocess_failures_return_json_errors(synthetic_web_client, monkeypatch, endpoint, label, failure):
    client, _ = synthetic_web_client

    def fake_run(cmd, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired(cmd, kwargs["timeout"])
        if failure == "invalid-json":
            return subprocess.CompletedProcess(cmd, 0, stdout="not JSON", stderr="")
        return subprocess.CompletedProcess(cmd, 1, stdout="", stderr="synthetic failure" if failure == "stderr" else "")

    monkeypatch.setattr(server.subprocess, "run", fake_run)
    response = client.post(
        f"/api/{endpoint}",
        json={"tournament_path": "2024/wimby2024", "gender": "male", "player1_name": "Synthetic Alpha", "player2_name": "Synthetic Beta"},
    )
    assert response.status_code == 500
    error = response.get_json()["error"]
    if failure == "invalid-json":
        assert error
    else:
        expected = {"stderr": "synthetic failure", "empty-stderr": f"{label} failed", "timeout": f"{label} timed out"}
        assert error == expected[failure]
